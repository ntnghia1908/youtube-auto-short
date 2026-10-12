"""SR1.2: importer S5 (.txt / .docx -> S2 text). Fixtures are self-written (no copyrighted text)."""

import io
import zipfile

import pytest

docx = pytest.importorskip("docx")

from auto_short.sutra import importer, markup  # noqa: E402
from auto_short.sutra.importer import ImportFileError, import_file  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def build_docx(fn) -> bytes:
    d = docx.Document()
    fn(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def add_footnotes(data: bytes, notes: dict[int, str]) -> bytes:
    """Add word/footnotes.xml (python-docx 1.2 has no footnote API)."""
    body = "".join(f'<w:footnote w:id="{i}"><w:p><w:r><w:t>{t}</w:t></w:r></w:p></w:footnote>' for i, t in notes.items())
    xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W}">'
           '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
           f'<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:t>SEP</w:t></w:r></w:p></w:footnote>{body}</w:footnotes>')
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            z.writestr(i, src.read(i.filename))
        z.writestr("word/footnotes.xml", xml)
    return out.getvalue()


def ref_run(paragraph, fid: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = OxmlElement("w:r")
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(fid))
    r.append(ref)
    paragraph._p.append(r)


# --- .txt ---------------------------------------------------------------------------------------------


def test_txt_plain_heuristics():
    raw = "Kinh Thử Nghiệm\n\n\n\nPHẨM MỘT\nDòng a\nDòng b\n\nChương 2 Hai\nDòng c\n".encode("utf-8")
    text, warnings = import_file("a.TXT", raw)
    assert text == "= Kinh Thử Nghiệm\n\n# PHẨM MỘT\nDòng a\nDòng b\n\n# Chương 2 Hai\nDòng c\n"
    assert warnings == []
    s = markup.parse(text)
    assert s.title == "Kinh Thử Nghiệm" and s.section_count == 2


def test_txt_bom_and_crlf():
    text, _ = import_file("a.txt", b"\xef\xbb\xbfKinh\r\n# Muc\r\nx\r\n")
    assert text.startswith("= Kinh\n# Muc\nx")
    assert "﻿" not in text and "\r" not in text


def test_txt_already_s2_unchanged():
    src = "= Kinh\n@ Nhóm\n# Mục\nDòng a[^1]\n[^1]: chú\n"
    text, _ = import_file("a.txt", src.encode("utf-8"))
    assert text == src


def test_txt_not_utf8_and_empty():
    with pytest.raises(ImportFileError) as e:
        import_file("a.txt", "Kinh Đại".encode("utf-16"))
    assert e.value.status == 422 and "UTF-8" in str(e.value)
    with pytest.raises(ImportFileError) as e:
        import_file("a.txt", "Kinh".encode("cp1258") + b"\xff")
    assert e.value.status == 422
    with pytest.raises(ImportFileError) as e:
        import_file("a.txt", b"\n  \n")
    assert e.value.status == 422


def test_txt_first_line_is_marker_adds_temp_title():
    text, warnings = import_file("a.txt", b"# Muc\nx\n")
    assert text.startswith("= B") and "# Muc" in text and warnings
    assert markup.parse(text).section_count == 1


def test_extension_and_size_limits():
    with pytest.raises(ImportFileError) as e:
        import_file("a.exe", b"x")
    assert e.value.status == 415
    with pytest.raises(ImportFileError) as e:
        import_file("noext", b"x")
    assert e.value.status == 415
    with pytest.raises(ImportFileError) as e:
        import_file("a.txt", b"a" * (importer.MAX_BYTES + 1))
    assert e.value.status == 413


# --- .docx --------------------------------------------------------------------------------------------


def test_docx_headings_paragraphs_softbreak_table():
    def fill(d):
        d.add_paragraph("Kinh Thử Docx")
        d.add_heading("Phẩm thứ nhất", level=1)
        d.add_paragraph("Dòng một")
        p = d.add_paragraph("Câu đầu")
        p.runs[0].add_break()
        p.add_run("Câu sau")
        d.add_paragraph("")
        d.add_paragraph("")
        d.add_paragraph("# không phải mục")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "ô A"
        t.cell(0, 1).text = "ô B"
        d.add_heading("Phẩm hai", level=2)
        d.add_paragraph("Dòng cuối")
    text, warnings = import_file("k.docx", build_docx(fill))
    assert text == ("= Kinh Thử Docx\n\n# Phẩm thứ nhất\nDòng một\nCâu đầu\nCâu sau\n\n\\# không phải mục\n"
                    "ô A\nô B\n\n# Phẩm hai\nDòng cuối\n")
    assert warnings == []
    s = markup.parse(text)
    assert [x.title for x in s.groups[0].sections] == ["Phẩm thứ nhất", "Phẩm hai"]
    assert [ln.text for ln in s.groups[0].sections[0].paragraphs[0]][:3] == ["Dòng một", "Câu đầu", "Câu sau"]


def test_docx_uppercase_heuristic_and_s2_passthrough():
    text, _ = import_file("k.docx", build_docx(lambda d: [d.add_paragraph(t) for t in ("Tên Kinh", "CHƯƠNG I", "nội dung")]))
    assert text == "= Tên Kinh\n\n# CHƯƠNG I\nnội dung\n"
    text, _ = import_file("k.docx", build_docx(lambda d: [d.add_paragraph(t) for t in ("= Tên", "# Mục", "dòng")]))
    assert text == "= Tên\n# Mục\ndòng\n"


def test_docx_footnotes_per_section():
    def fill(d):
        d.add_paragraph("Kinh Chú")
        d.add_heading("Mục một", level=1)
        p = d.add_paragraph("Có chú")
        ref_run(p, 1)
        p2 = d.add_paragraph("Thêm chú")
        ref_run(p2, 2)
        d.add_heading("Mục hai", level=1)
        p3 = d.add_paragraph("Lại có chú")
        ref_run(p3, 3)
    data = add_footnotes(build_docx(fill), {1: "chú một", 2: "chú hai", 3: "chú ba"})
    text, warnings = import_file("k.docx", data)
    assert text == ("= Kinh Chú\n\n# Mục một\nCó chú[^1]\nThêm chú[^2]\n\n[^1]: chú một\n[^2]: chú hai\n\n"
                    "# Mục hai\nLại có chú[^1]\n\n[^1]: chú ba\n")
    assert warnings == []
    s = markup.parse(text)
    assert s.groups[0].sections[1].notes == {"1": "chú ba"}


def test_docx_footnote_without_part_warns():
    def fill(d):
        d.add_paragraph("Kinh")
        ref_run(d.add_paragraph("x"), 1)
    text, warnings = import_file("k.docx", build_docx(fill))
    assert "[^" not in text and warnings


def test_docx_bad_zip_and_missing_document():
    for data in (b"not a zip", b"PK\x03\x04" + b"\0" * 10):
        with pytest.raises(ImportFileError) as e:
            import_file("k.docx", data)
        assert e.value.status == 422
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(ImportFileError) as e:
        import_file("k.docx", buf.getvalue())
    assert e.value.status == 422 and "document.xml" in str(e.value)


def test_docx_zip_limits_checked_before_python_docx(monkeypatch):
    monkeypatch.setattr(docx, "Document", lambda *a, **k: pytest.fail("python-docx must not open a rejected zip"))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<x/>")
        for i in range(importer.MAX_MEMBERS):
            z.writestr(f"f{i}", "")
    with pytest.raises(ImportFileError) as e:
        import_file("k.docx", buf.getvalue())
    assert e.value.status == 422 and "thành phần" in str(e.value)
    # total uncompressed size (a few MB compress to almost nothing)
    monkeypatch.setattr(importer, "MAX_UNZIPPED", 1024 * 1024)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<x/>")
        z.writestr("big.bin", b"\0" * (2 * 1024 * 1024))
    assert len(buf.getvalue()) < importer.MAX_BYTES
    with pytest.raises(ImportFileError) as e:
        import_file("k.docx", buf.getvalue())
    assert e.value.status == 422 and "50 MB" in str(e.value)


def test_docx_garbage_document_xml_is_422():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<not-closed")
    with pytest.raises(ImportFileError) as e:
        import_file("k.docx", buf.getvalue())
    assert e.value.status == 422


# --- SR1.5: text without any section heading ---------------------------------------------------------------


def test_docx_without_headings_gets_one_section_named_after_the_title():
    text, warnings = import_file("k.docx", build_docx(lambda d: [d.add_paragraph(t) for t in ("Tên Kinh", "dòng một", "dòng hai")]))
    assert text == "= Tên Kinh\n\n# Tên Kinh\n\ndòng một\ndòng hai\n"
    assert any("tiêu đề mục" in w for w in warnings)
    s = markup.parse(text)
    assert [x.title for x in s.groups[0].sections] == ["Tên Kinh"]


def test_txt_without_headings_and_s2_text_with_title_only():
    text, warnings = import_file("k.txt", "Bát Đại\nmột\n\nhai\n".encode())
    assert [x.title for x in markup.parse(text).groups[0].sections] == ["Bát Đại"]
    assert len(warnings) == 1
    text, _ = import_file("k.txt", "= Tên\n~ giới thiệu\nnội dung\n".encode())
    assert text == "= Tên\n~ giới thiệu\n\n# Tên\n\nnội dung\n"
    markup.parse(text)


def test_text_with_a_heading_or_without_content_is_unchanged():
    text, warnings = import_file("k.txt", "= Tên\n# Mục\nnội dung\n".encode())
    assert text == "= Tên\n# Mục\nnội dung\n" and warnings == []
    text, warnings = import_file("k.txt", "= Tên\n".encode())
    assert text == "= Tên\n" and warnings == []
    text, warnings = import_file("k.txt", "= Tên\n@ Nhóm\n# Mục\nnội dung\n".encode())
    assert "# Tên" not in text
