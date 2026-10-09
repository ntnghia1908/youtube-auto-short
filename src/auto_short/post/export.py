"""CP8.31 P16: export community posts as one Word ``.docx`` (python-docx, extra ``[web]``). Pure builder, no web
and no post-store dependency: the caller passes the ordered posts and the image library directory.

Canonical contract: docs/decisions/CP8.15-community-post-contract.md P16.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from . import images as post_images

IMAGE_WIDTH = Inches(5.5)  # fits the 6" text width of the default page


@dataclass
class ExportPost:
    title: str | None  # already upper-case (post.logic.text_parts)
    paragraphs: list[str]
    source: str | None  # P4 source line
    image: str | None = None  # library file name; a missing file is skipped


@dataclass
class ExportGroup:
    name: str  # "Shorts" | "Khai thị"
    posts: list[ExportPost] = field(default_factory=list)


@dataclass
class ExportEpisode:
    heading: str
    groups: list[ExportGroup] = field(default_factory=list)


def build_docx(path: Path, title: str, episodes: list[ExportEpisode], image_dir: Path, *, images: bool = True) -> None:
    """Title, then per episode: heading 1, per group heading 2, per post heading 3 (title line), image, paragraphs,
    source line (``images=False``: no picture at all). python-docx stores an image used by several posts once (it dedups by SHA-1)."""
    doc = Document()
    doc.add_heading(title, level=0)
    for episode in episodes:
        doc.add_heading(episode.heading, level=1)
        for group in episode.groups:
            doc.add_heading(group.name, level=2)
            for post in group.posts:
                if post.title:
                    doc.add_heading(post.title, level=3)
                image = post_images.resolve(image_dir, post.image) if (images and post.image) else None
                if image is not None:
                    try:
                        doc.add_picture(str(image), width=IMAGE_WIDTH)
                    except Exception:  # noqa: BLE001 - unreadable / unsupported image: keep the text
                        pass
                for text in post.paragraphs:
                    doc.add_paragraph(text)
                if post.source:
                    run = doc.add_paragraph().add_run(post.source)
                    run.italic = True
                    run.font.size = Pt(10)
    doc.save(str(path))
