"""Sutra import S5: ``.txt`` / ``.docx`` bytes -> sutra text S2 (+ warnings). No Auto Short imports.

The result is only a *suggestion*: the user edits it in the editor before saving. Contract:
docs/decisions/SR1-sutra-reader-contract.md S5 (PDF is S5a / SR1.3, not here).
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

from . import markup

MAX_BYTES = 5 * 1024 * 1024  # .txt / .docx upload
MAX_UNZIPPED = 50 * 1024 * 1024
MAX_MEMBERS = 2000
MAX_HEADING_CHARS = 80
EXTENSIONS = (".txt", ".docx")

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_FN_MARK = "\x00FN{}\x00"
_FN_RE = re.compile(r"\x00FN(-?\d+)\x00")
_HEADING_WORD = re.compile(r"^(CHƯƠNG|PHẨM|Chương|Phẩm)\s+\S")
_SPECIAL_START = re.compile(r"^([=~@#\\]|\*( |$)|\[\^\d{1,3}\]:)")
_MARKER_START = re.compile(r"^[=~@#]([ \t]|$)")


class ImportFileError(Exception):
    """Import failure; ``str(e)`` is a Vietnamese message, ``status`` the HTTP status (413 / 415 / 422)."""

    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.status = status


@dataclass
class _Para:
    text: str
    heading: bool = False


# --- entry point --------------------------------------------------------------------------------------


def import_file(name: str, data: bytes) -> tuple[str, list[str]]:
    """Convert an uploaded file to S2 text. Returns ``(text, warnings)``; raises :class:`ImportFileError`."""
    lower = (name or "").strip().lower()
    ext = next((e for e in EXTENSIONS if lower.endswith(e)), None)
    if ext is None:
        raise ImportFileError("chỉ nhận file .txt hoặc .docx", 415)
    if len(data) > MAX_BYTES:
        raise ImportFileError(f"file lớn hơn {MAX_BYTES // (1024 * 1024)} MB", 413)
    warnings: list[str] = []
    if ext == ".txt":
        paras = _read_txt(data)
        footnotes: dict[str, str] = {}
    else:
        paras, footnotes = _read_docx(data, warnings)
    text = _to_s2(paras, footnotes, warnings, docx=ext == ".docx")
    if not text.strip():
        raise ImportFileError("file không có chữ", 422)
    return text, warnings


# --- readers ------------------------------------------------------------------------------------------


def _read_txt(data: bytes) -> list[_Para]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise ImportFileError("file .txt không phải UTF-8 (hãy lưu lại ở dạng UTF-8)") from e
    return [_Para(ln.rstrip()) for ln in markup.normalize(text).split("\n")]


def _check_zip(data: bytes) -> zipfile.ZipFile:
    """Inspect the zip BEFORE python-docx touches it (member count, total size, document part)."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except (zipfile.BadZipFile, OSError, ValueError, EOFError) as e:
        raise ImportFileError("file .docx hỏng (không mở được)") from e
    if len(infos) > MAX_MEMBERS:
        raise ImportFileError(f"file .docx bất thường (quá {MAX_MEMBERS} thành phần)")
    if sum(i.file_size for i in infos) > MAX_UNZIPPED:
        raise ImportFileError("file .docx bất thường (giải nén quá 50 MB)")
    if "word/document.xml" not in {i.filename for i in infos}:
        raise ImportFileError("file .docx hỏng (thiếu word/document.xml)")
    return zf


def _xml_text(node, fn_ids: list[str] | None = None) -> str:
    """Text of a paragraph element in order: tabs -> space, soft breaks -> ``\\n``, footnote refs -> markers."""
    out: list[str] = []

    def walk(e) -> None:
        for ch in e:
            tag = ch.tag
            if tag == _W + "t":
                out.append(ch.text or "")
            elif tag == _W + "tab":
                out.append(" ")
            elif tag == _W + "cr" or (tag == _W + "br" and ch.get(_W + "type") in (None, "textWrapping")):
                out.append("\n")
            elif tag == _W + "footnoteReference":
                fid = ch.get(_W + "id")
                if fn_ids is not None and fid is not None:
                    out.append(_FN_MARK.format(fid))
            elif tag in (_W + "drawing", _W + "pict", _W + "object", _W + "del", _W + "instrText"):
                continue  # images / text boxes / deleted text / field codes
            else:
                walk(ch)

    walk(node)
    return "".join(out)


def _read_docx(data: bytes, warnings: list[str]) -> tuple[list[_Para], dict[str, str]]:
    zf = _check_zip(data)
    try:
        import docx  # python-docx, extra [web]
        from docx.text.paragraph import Paragraph

        doc = docx.Document(io.BytesIO(data))
        body = doc.element.body
    except ImportError as e:  # pragma: no cover - dependency missing
        raise ImportFileError("máy chủ chưa cài python-docx, không đọc được .docx") from e
    except Exception as e:  # noqa: BLE001 - python-docx / lxml raise many types on broken input
        raise ImportFileError("file .docx hỏng (không đọc được nội dung)") from e

    paras: list[_Para] = []
    has_fn = False

    def style_is_heading(p_el) -> bool:
        try:
            name = (Paragraph(p_el, doc).style.name or "").lower()
        except Exception:  # noqa: BLE001
            return False
        return name.startswith("heading") or name == "title"

    def add_paragraph(p_el, *, allow_heading: bool) -> None:
        nonlocal has_fn
        raw = _xml_text(p_el, [])
        if _FN_RE.search(raw):
            has_fn = True
        lines = [ln.strip() for ln in markup.normalize(raw).split("\n")]
        if allow_heading and style_is_heading(p_el) and any(lines):
            paras.append(_Para(" ".join(ln for ln in lines if ln), True))
        else:
            paras.extend(_Para(ln) for ln in lines)

    def walk_block(parent, *, in_table: bool) -> None:
        for ch in parent:
            if ch.tag == _W + "p":
                add_paragraph(ch, allow_heading=not in_table)
            elif ch.tag == _W + "tbl":
                if in_table:
                    continue  # nested tables: skipped
                for tr in ch.iter(_W + "tr"):
                    for tc in tr.findall(_W + "tc"):
                        walk_block(tc, in_table=True)
            elif ch.tag == _W + "sdt":
                for content in ch.findall(_W + "sdtContent"):
                    walk_block(content, in_table=in_table)

    walk_block(body, in_table=False)
    footnotes = _read_footnotes(zf, warnings) if has_fn else {}
    return paras, footnotes


def _read_footnotes(zf: zipfile.ZipFile, warnings: list[str]) -> dict[str, str]:
    """``word/footnotes.xml`` -> ``{footnote id: text}`` (separators skipped)."""
    names = {i.filename for i in zf.infolist()}
    if "word/footnotes.xml" not in names:
        warnings.append("không đọc được chú thích Word (thiếu word/footnotes.xml); dấu chú thích đã bị bỏ")
        return {}
    try:
        from lxml import etree

        with zf.open("word/footnotes.xml") as fh:
            raw = fh.read(MAX_UNZIPPED + 1)
        root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))
    except Exception:  # noqa: BLE001
        warnings.append("không đọc được chú thích Word (footnotes.xml hỏng); dấu chú thích đã bị bỏ")
        return {}
    out: dict[str, str] = {}
    for fn in root.findall(_W + "footnote"):
        if fn.get(_W + "type") in ("separator", "continuationSeparator", "continuationNotice"):
            continue
        fid = fn.get(_W + "id")
        parts = []
        for p in fn.iter(_W + "p"):
            t = " ".join(markup.normalize(_xml_text(p)).split())
            if t:
                parts.append(t)
        if fid is not None and parts:
            out[fid] = " ".join(parts)
    return out


# --- conversion to S2 ---------------------------------------------------------------------------------


def _is_heading_line(text: str) -> bool:
    t = text.strip()
    if not t or len(t) > MAX_HEADING_CHARS or _MARKER_START.match(t):
        return False
    if _HEADING_WORD.match(t):
        return True
    letters = [c for c in t if c.isalpha()]
    return len(letters) >= 3 and t.isupper()


def _escape(text: str) -> str:
    return "\\" + text if _SPECIAL_START.match(text) else text


def _to_s2(paras: list[_Para], footnotes: dict[str, str], warnings: list[str], *, docx: bool) -> str:
    # drop leading / trailing blanks, collapse runs of blank lines
    lines: list[_Para] = []
    for p in paras:
        if not p.text.strip():
            if lines and lines[-1].text:
                lines.append(_Para(""))
            continue
        lines.append(p)
    while lines and not lines[-1].text:
        lines.pop()
    if not lines:
        return ""

    already = bool(re.match(r"^=([ \t]|$)", lines[0].text))
    out: list[_Para] = []
    if already:
        out = [_Para(p.text) for p in lines]  # S2 text: unchanged (docx heading styles are not re-marked)
    else:
        first = lines[0].text
        if _MARKER_START.match(first):
            warnings.append("văn bản chưa có dòng '= <tên>'; đã thêm tên tạm, hãy sửa lại")
            out.append(_Para("= Bộ kinh mới"))
            out.append(_Para(""))
            title_done = True
        else:
            title_done = False
        for i, p in enumerate(lines):
            text = p.text
            if not title_done and i == 0:
                out.append(_Para("= " + text))
                continue
            if p.heading or _is_heading_line(text):
                out.append(_Para(""))
                out.append(_Para("# " + text))
            else:
                out.append(_Para(_escape(text) if docx else text))

    return _finish(out, footnotes, warnings)


def _finish(lines: list[_Para], footnotes: dict[str, str], warnings: list[str]) -> str:
    """Resolve footnote markers to per-section ``[^n]`` + definitions at the end of the section; join."""
    out: list[str] = []
    defs: list[str] = []
    counter = 0

    def flush() -> None:
        nonlocal defs, counter
        if defs:
            while out and out[-1] == "":
                out.pop()
            out.append("")
            out.extend(defs)
            out.append("")
        defs, counter = [], 0

    def repl(m: re.Match) -> str:
        nonlocal counter
        fid = m.group(1)
        note = footnotes.get(fid)
        if note is None:
            warnings.append(f"chú thích Word #{fid} không có nội dung; đã bỏ dấu chú thích")
            return ""
        if counter >= 999:
            warnings.append("quá 999 chú thích trong một mục; đã bỏ dấu chú thích thừa")
            return ""
        counter += 1
        defs.append(f"[^{counter}]: {note}")
        return f"[^{counter}]"

    for p in lines:
        text = p.text
        if text.startswith("#") and re.match(r"^#([ \t]|$)", text):
            flush()
        if "\x00" in text:
            text = _FN_RE.sub(repl, text)
        out.append(text.rstrip())
    flush()
    # collapse blank runs once more (heading insertion / removed markers can leave doubles)
    final: list[str] = []
    for ln in out:
        if ln == "" and (not final or final[-1] == ""):
            continue
        final.append(ln)
    while final and final[-1] == "":
        final.pop()
    return "\n".join(final) + "\n"
